from collections.abc import Callable
from pathlib import Path

import docx
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

from paddlepdf.pdftools.contract import ToolSpec
from paddlepdf.pdftools.office import TOOLS
from paddlepdf.pdftools.runner import run_tool

type MakePdf = Callable[..., Path]

SPECS = {spec.name: spec for spec in TOOLS}
TABLE = [["Name", "Qty", "Price"], ["Apple", "3", "1.50"], ["Banana", "12", "0.25"]]


def _spec(name: str) -> ToolSpec:
    return SPECS[name]


def _table_pdf(path: Path) -> Path:
    canvas = Canvas(str(path), pagesize=letter)
    canvas.setFont("Helvetica", 12)
    for row_index, row in enumerate(TABLE):
        for column_index, value in enumerate(row):
            canvas.drawString(72 + 150 * column_index, 700 - 20 * row_index, value)
    canvas.showPage()
    canvas.drawString(72, 700, "Closing remarks only")
    canvas.showPage()
    canvas.save()
    return path


def _image_pdf(path: Path) -> Path:
    canvas = Canvas(str(path), pagesize=letter)
    canvas.drawString(72, 720, "Figure page")
    canvas.drawImage(ImageReader(Image.new("RGB", (60, 40), "red")), 72, 500, 120, 80)
    canvas.showPage()
    canvas.drawString(72, 720, "Second page")
    canvas.showPage()
    canvas.save()
    return path


def test_lane_registers_every_paperpilot_tool() -> None:
    assert sorted(SPECS) == sorted(
        [
            "pdf_to_docx",
            "pdf_to_xlsx",
            "pdf_to_pptx",
            "office_to_markdown",
            "office_extract_text",
            "office_modernize_legacy",
            "excel_to_json",
            "excel_to_csv",
        ]
    )


def test_pdf_to_docx_writes_text_page_breaks_and_images(tmp_path: Path) -> None:
    source = _image_pdf(tmp_path / "figure.pdf")
    target = tmp_path / "out" / "figure.docx"

    result = run_tool(
        _spec("pdf_to_docx"), {"input_file": str(source), "output_file": str(target)}
    )

    assert result.status == "success", result.errors
    assert result.output_files == (str(target),)
    document = docx.Document(str(target))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "Figure page" in texts
    assert "Second page" in texts
    assert document.element.xml.count('w:type="page"') == 1
    assert len(document.inline_shapes) == 1
    assert result.data == {"pages": 2, "paragraphs": 2, "images": 1}


def test_pdf_to_docx_without_images(make_pdf: MakePdf, tmp_path: Path) -> None:
    source = make_pdf(pages=3, text="Body line")
    target = tmp_path / "plain.docx"

    result = run_tool(
        _spec("pdf_to_docx"),
        {
            "input_file": str(source),
            "output_file": str(target),
            "include_images": False,
        },
    )

    assert result.status == "success", result.errors
    document = docx.Document(str(target))
    texts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert texts == [
        "Page 1",
        "Body line",
        "Page 2",
        "Body line",
        "Page 3",
        "Body line",
    ]
    assert document.element.xml.count('w:type="page"') == 2
    assert len(document.inline_shapes) == 0


def test_pdf_to_xlsx_splits_table_columns_by_word_position(tmp_path: Path) -> None:
    source = _table_pdf(tmp_path / "table.pdf")
    target = tmp_path / "table.xlsx"

    result = run_tool(
        _spec("pdf_to_xlsx"), {"input_file": str(source), "output_file": str(target)}
    )

    assert result.status == "success", result.errors
    workbook = load_workbook(target)
    assert workbook.sheetnames == ["Page 1", "Page 2"]
    first = workbook["Page 1"]
    second = workbook["Page 2"]
    assert isinstance(first, Worksheet)
    assert isinstance(second, Worksheet)
    assert [list(row) for row in first.iter_rows(values_only=True)] == TABLE
    assert [list(row) for row in second.iter_rows(values_only=True)] == [
        ["Closing remarks only"]
    ]
    assert result.data == {
        "pages": [
            {"page": 1, "table_detected": True, "rows": 3, "columns": 3},
            {"page": 2, "table_detected": False, "rows": 1, "columns": 1},
        ]
    }


def test_pdf_to_pptx_renders_one_picture_slide_per_page(
    make_pdf: MakePdf, tmp_path: Path
) -> None:
    source = make_pdf(pages=2)
    target = tmp_path / "deck.pptx"

    result = run_tool(
        _spec("pdf_to_pptx"),
        {"input_file": str(source), "output_file": str(target), "dpi": 72},
    )

    assert result.status == "success", result.errors
    presentation = Presentation(str(target))
    assert presentation.slide_width == 612 * 914_400 // 72
    assert presentation.slide_height == 792 * 914_400 // 72
    slides = list(presentation.slides)
    assert len(slides) == 2
    for number, slide in enumerate(slides, start=1):
        kinds = [shape.shape_type for shape in slide.shapes]
        assert kinds == [MSO_SHAPE_TYPE.PICTURE]
        frame = slide.notes_slide.notes_text_frame
        assert frame is not None
        assert frame.text == f"Page {number}"


def test_pdf_tools_fail_cleanly_for_missing_or_wrong_inputs(tmp_path: Path) -> None:
    missing = run_tool(
        _spec("pdf_to_pptx"),
        {
            "input_file": str(tmp_path / "nope.pdf"),
            "output_file": str(tmp_path / "x.pptx"),
        },
    )
    bad_dpi = run_tool(
        _spec("pdf_to_pptx"),
        {"input_file": str(tmp_path / "nope.pdf"), "output_file": "x.pptx", "dpi": 5},
    )

    assert missing.status == "failed"
    assert missing.errors
    assert "not found" in missing.errors[0]
    assert bad_dpi.status == "failed"
    assert bad_dpi.errors
