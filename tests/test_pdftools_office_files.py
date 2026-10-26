import json
from pathlib import Path
from typing import TYPE_CHECKING, cast

import docx
from openpyxl import Workbook
from pptx import Presentation
from pptx.shapes.autoshape import Shape

from paddlepdf.pdftools.contract import ToolSpec
from paddlepdf.pdftools.office import TOOLS
from paddlepdf.pdftools.runner import run_tool

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

SPECS = {spec.name: spec for spec in TOOLS}


def _spec(name: str) -> ToolSpec:
    return SPECS[name]


def _word(path: Path) -> Path:
    document = docx.Document()
    document.add_heading("Report", level=1)
    paragraph = document.add_paragraph("Plain and ")
    paragraph.add_run("bold").bold = True
    document.add_paragraph("item one", style="List Bullet")
    document.add_paragraph("item two", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    for row_index, values in enumerate([("A", "B"), ("1", "2")]):
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    document.save(str(path))
    return path


def _deck(path: Path) -> Path:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    title = slide.shapes.title
    assert title is not None
    title.text = "Deck title"
    body = slide.placeholders[1]
    assert isinstance(body, Shape)
    body.text_frame.text = "First point"
    frame = slide.notes_slide.notes_text_frame
    assert frame is not None
    frame.text = "Speaker note"
    presentation.save(str(path))
    return path


def _workbook(path: Path) -> Path:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Sheet1"
    first.append(["Header1", "Header2"])
    first.append(["Value1", 42])
    second = cast("Worksheet", workbook.create_sheet("Sheet 2"))
    second.append(["OnlyOne"])
    workbook.save(path)
    return path


def test_office_to_markdown_docx_structure(tmp_path: Path) -> None:
    target = tmp_path / "report.md"

    result = run_tool(
        _spec("office_to_markdown"),
        {"input_file": str(_word(tmp_path / "r.docx")), "output_file": str(target)},
    )

    assert result.status == "success", result.errors
    assert target.read_text(encoding="utf-8") == (
        "# Report\n\nPlain and **bold**\n\n- item one\n- item two\n\n"
        "| A | B |\n| --- | --- |\n| 1 | 2 |\n"
    )


def test_office_to_markdown_pptx_and_xlsx(tmp_path: Path) -> None:
    deck_md = tmp_path / "deck.md"
    sheet_md = tmp_path / "sheet.md"

    deck = run_tool(
        _spec("office_to_markdown"),
        {"input_file": str(_deck(tmp_path / "d.pptx")), "output_file": str(deck_md)},
    )
    sheet = run_tool(
        _spec("office_to_markdown"),
        {
            "input_file": str(_workbook(tmp_path / "w.xlsx")),
            "output_file": str(sheet_md),
        },
    )

    assert deck.status == "success", deck.errors
    assert deck_md.read_text(encoding="utf-8") == (
        "## Slide 1\n\n### Deck title\n\nFirst point\n\n*Notes:* Speaker note\n"
    )
    assert sheet.status == "success", sheet.errors
    assert sheet_md.read_text(encoding="utf-8") == (
        "## Sheet1\n\n| Header1 | Header2 |\n| --- | --- |\n| Value1 | 42 |\n\n"
        "## Sheet 2\n\n| OnlyOne |\n| --- |\n"
    )


def test_office_extract_text_tabs_table_cells(tmp_path: Path) -> None:
    target = tmp_path / "report.txt"

    result = run_tool(
        _spec("office_extract_text"),
        {"input_file": str(_word(tmp_path / "r.docx")), "output_file": str(target)},
    )

    assert result.status == "success", result.errors
    assert target.read_text(encoding="utf-8") == (
        "Report\nPlain and bold\nitem one\nitem two\nA\tB\n1\t2\n"
    )
    assert result.data == {"format": "docx", "blocks": 5, "characters": 48}


def test_office_modernize_legacy_named_ooxml(tmp_path: Path) -> None:
    legacy = tmp_path / "old.doc"
    legacy.write_bytes(_word(tmp_path / "source.docx").read_bytes())
    target = tmp_path / "old.docx"

    result = run_tool(
        _spec("office_modernize_legacy"),
        {"input_file": str(legacy), "output_file": str(target)},
    )

    assert result.status == "success", result.errors
    assert docx.Document(str(target)).paragraphs[0].text == "Report"
    assert result.data == {"format": "docx", "source_container": "ooxml"}


def test_office_modernize_legacy_rejects_ole2_binaries(tmp_path: Path) -> None:
    legacy = tmp_path / "binary.xls"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504))

    result = run_tool(
        _spec("office_modernize_legacy"),
        {"input_file": str(legacy), "output_file": str(tmp_path / "binary.xlsx")},
    )

    assert result.status == "failed"
    assert result.errors
    assert "legacy binary Office file" in result.errors[0]
    assert not (tmp_path / "binary.xlsx").exists()


def test_excel_to_json_with_headers_and_file(tmp_path: Path) -> None:
    target = tmp_path / "book.json"

    result = run_tool(
        _spec("excel_to_json"),
        {"input_file": str(_workbook(tmp_path / "w.xlsx")), "output_file": str(target)},
    )

    assert result.status == "success", result.errors
    expected = {
        "sheets": [
            {
                "name": "Sheet1",
                "rows_count": 2,
                "columns": ["Header1", "Header2"],
                "data": [{"Header1": "Value1", "Header2": 42}],
            },
            {"name": "Sheet 2", "rows_count": 1, "columns": ["OnlyOne"], "data": []},
        ]
    }
    assert result.data == expected
    assert json.loads(target.read_text(encoding="utf-8")) == expected


def test_excel_to_json_single_sheet_rows_and_missing_sheet(tmp_path: Path) -> None:
    book = _workbook(tmp_path / "w.xlsx")

    raw = run_tool(
        _spec("excel_to_json"),
        {"input_file": str(book), "sheet": "Sheet1", "has_headers": False},
    )
    missing = run_tool(
        _spec("excel_to_json"), {"input_file": str(book), "sheet": "Nope"}
    )

    assert raw.status == "success", raw.errors
    assert raw.output_files == ()
    assert raw.data == {
        "sheets": [
            {
                "name": "Sheet1",
                "rows_count": 2,
                "columns": [],
                "data": [["Header1", "Header2"], ["Value1", 42]],
            }
        ]
    }
    assert missing.status == "failed"
    assert missing.errors
    assert "Sheet not found: Nope" in missing.errors[0]


def test_excel_to_csv_named_sheet(tmp_path: Path) -> None:
    target = tmp_path / "one.csv"

    result = run_tool(
        _spec("excel_to_csv"),
        {
            "input_file": str(_workbook(tmp_path / "w.xlsx")),
            "output_file": str(target),
            "sheet": "Sheet1",
        },
    )

    assert result.status == "success", result.errors
    assert result.output_files == (str(target),)
    assert target.read_bytes() == b"Header1,Header2\r\nValue1,42\r\n"


def test_excel_to_csv_multi_sheet_writes_one_file_per_sheet(tmp_path: Path) -> None:
    target = tmp_path / "report.csv"

    result = run_tool(
        _spec("excel_to_csv"),
        {"input_file": str(_workbook(tmp_path / "w.xlsx")), "output_file": str(target)},
    )

    assert result.status == "success", result.errors
    first = tmp_path / "report_Sheet1.csv"
    second = tmp_path / "report_Sheet_2.csv"
    assert result.output_files == (str(target), str(first), str(second))
    assert (
        target.read_bytes() == first.read_bytes() == b"Header1,Header2\r\nValue1,42\r\n"
    )
    assert second.read_bytes() == b"OnlyOne\r\n"
