"""pdf_to_docx (page text + embedded images) and pdf_to_pptx (rendered page slides)."""

from collections.abc import Iterator, Mapping
from io import BytesIO
from typing import Annotated, Final, cast

import docx
import pikepdf
from docx.document import Document as DocxDocument
from docx.shared import Emu, Inches
from PIL import Image
from pptx import Presentation
from pptx.util import Emu as SlideEmu
from pydantic import Field

from paddlepdf.pdftools.common import input_path, open_pdf, output_path
from paddlepdf.pdftools.contract import ToolOutcome
from paddlepdf.pdftools.office.formats import page_lines, xml_safe
from paddlepdf.pdftools.pdfium import page_count, page_texts, render_page

_EMU_PER_INCH: Final = 914_400
_IMAGE_DPI: Final = 150
_MAX_IMAGE_WIDTH_EMU: Final = 4_500_000
_MAX_SLIDE_EMU: Final = 51_206_400
_MIN_SLIDE_EMU: Final = _EMU_PER_INCH
_PNG_MODES: Final = frozenset({"1", "L", "LA", "P", "RGB", "RGBA"})
_IMAGE_ERRORS: Final = (
    pikepdf.UnsupportedImageTypeError,
    NotImplementedError,
    pikepdf.PdfError,
    OSError,
    ValueError,
)

PasswordParam = Annotated[
    str | None, Field(description="Password for an encrypted input PDF.")
]


def _png(image: Image.Image) -> BytesIO:
    buffer = BytesIO()
    (image if image.mode in _PNG_MODES else image.convert("RGB")).save(buffer, "PNG")
    buffer.seek(0)
    return buffer


def _image_streams(images: object) -> Mapping[str, pikepdf.Stream]:
    return cast("Mapping[str, pikepdf.Stream]", images)


def _page_images(
    page: pikepdf.Page, page_number: int, warnings: list[str]
) -> Iterator[Image.Image]:
    for name, stream in _image_streams(page.get_images()).items():
        try:
            yield pikepdf.PdfImage(stream).as_pil_image()
        except _IMAGE_ERRORS as error:
            warnings.append(f"Page {page_number}: skipped image {name}: {error}")


def _add_picture(document: DocxDocument, image: Image.Image) -> None:
    width = image.width * _EMU_PER_INCH // _IMAGE_DPI
    height = image.height * _EMU_PER_INCH // _IMAGE_DPI
    if width > _MAX_IMAGE_WIDTH_EMU:
        height = height * _MAX_IMAGE_WIDTH_EMU // width
        width = _MAX_IMAGE_WIDTH_EMU
    document.add_picture(_png(image), width=Emu(width), height=Emu(max(height, 1)))


def pdf_to_docx(
    input_file: Annotated[str, Field(description="Source PDF.")],
    output_file: Annotated[str, Field(description="Destination .docx path.")],
    password: PasswordParam = None,
    *,
    include_images: Annotated[
        bool, Field(description="Embed each page's image XObjects after its text.")
    ] = True,
) -> ToolOutcome:
    source = input_path(input_file)
    target = output_path(output_file, source)
    texts = page_texts(source, password)
    document = docx.Document()
    for section in document.sections:
        margin = Inches(0.5)
        section.top_margin = section.bottom_margin = margin
        section.left_margin = section.right_margin = margin
    warnings: list[str] = []
    paragraphs = images = 0
    with open_pdf(source, password) as pdf:
        for index, (page, text) in enumerate(zip(pdf.pages, texts, strict=True)):
            if index:
                document.add_page_break()
            for line in page_lines(text):
                document.add_paragraph(line)
                paragraphs += 1
            if include_images:
                for image in _page_images(page, index + 1, warnings):
                    _add_picture(document, image)
                    images += 1
    document.save(str(target))
    if paragraphs == 0:
        warnings.append("No extractable text found; scanned PDFs need OCR first")
    return ToolOutcome(
        output_files=(str(target),),
        data={"pages": len(texts), "paragraphs": paragraphs, "images": images},
        warnings=tuple(warnings),
    )


def _slide_size(image: Image.Image, dpi: float) -> tuple[int, int]:
    width = image.width * _EMU_PER_INCH / dpi
    height = image.height * _EMU_PER_INCH / dpi
    shrink = min(1.0, _MAX_SLIDE_EMU / max(width, height))
    return (
        max(int(width * shrink), _MIN_SLIDE_EMU),
        max(int(height * shrink), _MIN_SLIDE_EMU),
    )


def pdf_to_pptx(
    input_file: Annotated[str, Field(description="Source PDF.")],
    output_file: Annotated[str, Field(description="Destination .pptx path.")],
    dpi: Annotated[
        int, Field(ge=36, le=600, description="Render resolution for slide images.")
    ] = 150,
    password: PasswordParam = None,
) -> ToolOutcome:
    source = input_path(input_file)
    target = output_path(output_file, source)
    count = page_count(source, password)
    texts = page_texts(source, password)
    presentation = Presentation()
    first = render_page(source, 0, dpi, password)
    slide_width, slide_height = _slide_size(first, dpi)
    presentation.slide_width = SlideEmu(slide_width)
    presentation.slide_height = SlideEmu(slide_height)
    layout = presentation.slide_layouts[6]
    for index in range(count):
        image = first if index == 0 else render_page(source, index, dpi, password)
        slide = presentation.slides.add_slide(layout)
        scale = min(slide_width / image.width, slide_height / image.height)
        width, height = int(image.width * scale), int(image.height * scale)
        slide.shapes.add_picture(
            _png(image),
            SlideEmu((slide_width - width) // 2),
            SlideEmu((slide_height - height) // 2),
            SlideEmu(width),
            SlideEmu(height),
        )
        notes = xml_safe(texts[index].replace("\r\n", "\n")).strip()
        frame = slide.notes_slide.notes_text_frame if notes else None
        if frame is not None:
            frame.text = notes
    presentation.save(str(target))
    return ToolOutcome(
        output_files=(str(target),),
        data={
            "slides": count,
            "dpi": dpi,
            "slide_size_emu": [slide_width, slide_height],
        },
    )
