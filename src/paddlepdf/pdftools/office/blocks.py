"""Read .docx/.pptx/.xlsx into ordered content blocks, then render Markdown or text."""

from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path

from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.slide import Slide

from paddlepdf.pdftools.office.formats import (
    OfficeFamily,
    cell_text,
    open_docx,
    open_pptx,
    open_xlsx,
    sheet_rows,
    worksheets,
)


class BlockKind(StrEnum):
    HEADING = "heading"
    PARAGRAPH = "paragraph"
    BULLET = "bullet"
    NUMBERED = "numbered"
    TABLE = "table"


@dataclass(frozen=True, slots=True)
class Block:
    kind: BlockKind
    text: str = ""
    rich: str = ""
    level: int = 0
    rows: tuple[tuple[str, ...], ...] = ()


def _heading_level(style: str) -> int:
    if style == "Title":
        return 1
    suffix = style.removeprefix("Heading ")
    return int(suffix) if suffix != style and suffix.isdigit() else 0


def _rich_runs(paragraph: DocxParagraph) -> str:
    parts: list[str] = []
    for run in paragraph.runs:
        text = run.text
        if not text.strip():
            parts.append(text)
            continue
        marker = ("**" if run.bold else "") + ("*" if run.italic else "")
        core = text.strip()
        lead = text[: len(text) - len(text.lstrip())]
        trail = text[len(text.rstrip()) :]
        parts.append(f"{lead}{marker}{core}{marker[::-1]}{trail}")
    return "".join(parts).strip()


def _docx_paragraph(paragraph: DocxParagraph) -> Block | None:
    text = paragraph.text.strip()
    if not text:
        return None
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    level = _heading_level(style)
    if level:
        return Block(BlockKind.HEADING, text, text, level)
    rich = _rich_runs(paragraph) or text
    if style.startswith("List Number"):
        return Block(BlockKind.NUMBERED, text, rich)
    if style.startswith("List"):
        return Block(BlockKind.BULLET, text, rich)
    return Block(BlockKind.PARAGRAPH, text, rich)


def _table_block(rows: Iterable[Iterable[str]]) -> Block | None:
    grid = tuple(tuple(cell.strip() for cell in row) for row in rows)
    if not any(any(row) for row in grid):
        return None
    return Block(BlockKind.TABLE, rows=grid)


def _docx_blocks(path: Path) -> Iterator[Block]:
    document = open_docx(path)
    for item in document.iter_inner_content():
        if isinstance(item, DocxTable):
            block = _table_block([cell.text for cell in row.cells] for row in item.rows)
        else:
            block = _docx_paragraph(item)
        if block is not None:
            yield block


def _text_shape_blocks(shape: Shape, title_id: int | None) -> Iterator[Block]:
    if shape.shape_id == title_id:
        title = shape.text_frame.text.strip()
        if title:
            yield Block(BlockKind.HEADING, title, title, 3)
        return
    for paragraph in shape.text_frame.paragraphs:
        text = paragraph.text.strip()
        if text:
            kind = BlockKind.BULLET if paragraph.level > 0 else BlockKind.PARAGRAPH
            yield Block(kind, text, text)


def _shape_blocks(shape: BaseShape, title_id: int | None) -> Iterator[Block]:
    if isinstance(shape, GroupShape):
        for child in shape.shapes:
            yield from _shape_blocks(child, title_id)
    elif isinstance(shape, GraphicFrame) and shape.has_table:
        block = _table_block(
            [cell.text for cell in row.cells] for row in shape.table.rows
        )
        yield from (block,) if block is not None else ()
    elif isinstance(shape, Shape) and shape.has_text_frame:
        yield from _text_shape_blocks(shape, title_id)


def _slide_blocks(slide: Slide, number: int) -> Iterator[Block]:
    yield Block(BlockKind.HEADING, f"Slide {number}", f"Slide {number}", 2)
    title = slide.shapes.title
    title_id = title.shape_id if title is not None else None
    for shape in slide.shapes:
        yield from _shape_blocks(shape, title_id)
    if slide.has_notes_slide:
        frame = slide.notes_slide.notes_text_frame
        notes = frame.text.strip() if frame is not None else ""
        if notes:
            yield Block(BlockKind.PARAGRAPH, f"Notes: {notes}", f"*Notes:* {notes}")


def _pptx_blocks(path: Path) -> Iterator[Block]:
    presentation = open_pptx(path)
    for number, slide in enumerate(presentation.slides, start=1):
        yield from _slide_blocks(slide, number)


def _xlsx_blocks(path: Path) -> Iterator[Block]:
    workbook = open_xlsx(path)
    for sheet in worksheets(workbook):
        yield Block(BlockKind.HEADING, sheet.title, sheet.title, 2)
        block = _table_block(
            [cell_text(value) for value in row] for row in sheet_rows(sheet)
        )
        if block is not None:
            yield block
    workbook.close()


def read_blocks(path: Path, family: OfficeFamily) -> list[Block]:
    readers = {"docx": _docx_blocks, "pptx": _pptx_blocks, "xlsx": _xlsx_blocks}
    return list(readers[family](path))


def _markdown_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def _markdown_table(rows: tuple[tuple[str, ...], ...]) -> str:
    width = max(len(row) for row in rows)
    padded = [[*row, *([""] * (width - len(row)))] for row in rows]
    lines = ["| " + " | ".join(_markdown_cell(cell) for cell in padded[0]) + " |"]
    lines.append("| " + " | ".join(["---"] * width) + " |")
    lines.extend(
        "| " + " | ".join(_markdown_cell(cell) for cell in row) + " |"
        for row in padded[1:]
    )
    return "\n".join(lines)


def to_markdown(blocks: list[Block]) -> str:
    parts: list[str] = []
    previous: BlockKind | None = None
    for block in blocks:
        match block.kind:
            case BlockKind.HEADING:
                chunk = f"{'#' * min(block.level, 6)} {block.text}"
            case BlockKind.BULLET:
                chunk = f"- {block.rich}"
            case BlockKind.NUMBERED:
                chunk = f"1. {block.rich}"
            case BlockKind.TABLE:
                chunk = _markdown_table(block.rows)
            case BlockKind.PARAGRAPH:
                chunk = block.rich
        same_list = block.kind == previous and block.kind in {
            BlockKind.BULLET,
            BlockKind.NUMBERED,
        }
        parts.append(("\n" if same_list else "\n\n") + chunk if parts else chunk)
        previous = block.kind
    return "".join(parts) + "\n" if parts else ""


def to_text(blocks: list[Block]) -> str:
    lines: list[str] = []
    for block in blocks:
        if block.kind is BlockKind.TABLE:
            lines.extend("\t".join(row) for row in block.rows)
        else:
            lines.append(block.text)
    return "\n".join(lines) + "\n" if lines else ""
