"""Office package opening, legacy-binary detection, and XML-safe text helpers."""

import datetime as dt
import re
import zipfile
from pathlib import Path
from typing import Final, Literal

import docx
import pptx
from docx.document import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError as DocxPackageError
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet
from pptx.exc import PackageNotFoundError as PptxPackageError
from pptx.presentation import Presentation

from paddlepdf.pdftools.contract import PdfToolError

type OfficeFamily = Literal["docx", "xlsx", "pptx"]

OLE_MAGIC: Final = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
LEGACY_SUFFIXES: Final = (".doc", ".xls", ".ppt")
_FAMILY_PARTS: Final[tuple[tuple[str, OfficeFamily], ...]] = (
    ("word/document.xml", "docx"),
    ("xl/workbook.xml", "xlsx"),
    ("ppt/presentation.xml", "pptx"),
)
_XML_ILLEGAL: Final = re.compile(
    "[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]"
)
_OPEN_ERRORS: Final = (
    zipfile.BadZipFile,
    KeyError,
    ValueError,
    DocxPackageError,
    PptxPackageError,
    InvalidFileException,
)


def xml_safe(text: str) -> str:
    """Drop characters that OOXML writers reject (control chars, lone surrogates)."""
    return _XML_ILLEGAL.sub("", text)


def page_lines(text: str) -> list[str]:
    return [clean for line in text.splitlines() if (clean := xml_safe(line).strip())]


def reject_legacy_binary(path: Path) -> None:
    with path.open("rb") as handle:
        head = handle.read(len(OLE_MAGIC))
    if head == OLE_MAGIC:
        msg = (
            f"{path.name} is a legacy binary Office file (OLE2 .doc/.xls/.ppt). "
            "The bundled Office libraries only read Office Open XML; re-save it as "
            ".docx/.xlsx/.pptx (for example `soffice --headless --convert-to docx "
            f"{path.name}`) and retry."
        )
        raise PdfToolError(msg)


def detect_family(path: Path) -> OfficeFamily:
    reject_legacy_binary(path)
    try:
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
    except zipfile.BadZipFile as error:
        msg = f"{path.name} is not an Office Open XML package (.docx/.xlsx/.pptx)"
        raise PdfToolError(msg) from error
    for part, family in _FAMILY_PARTS:
        if part in names:
            return family
    msg = f"{path.name} is a ZIP file but not a Word, Excel, or PowerPoint package"
    raise PdfToolError(msg)


def _open_failed(path: Path, error: Exception) -> PdfToolError:
    return PdfToolError(f"Failed to open Office file {path}: {error}")


def open_docx(path: Path) -> DocxDocument:
    reject_legacy_binary(path)
    try:
        return docx.Document(str(path))
    except _OPEN_ERRORS as error:
        raise _open_failed(path, error) from error


def open_pptx(path: Path) -> Presentation:
    reject_legacy_binary(path)
    try:
        return pptx.Presentation(str(path))
    except _OPEN_ERRORS as error:
        raise _open_failed(path, error) from error


def open_xlsx(path: Path) -> Workbook:
    """Open through a handle so a renamed workbook is not refused for its suffix."""
    reject_legacy_binary(path)
    try:
        with path.open("rb") as handle:
            return load_workbook(handle, data_only=True)
    except _OPEN_ERRORS as error:
        raise _open_failed(path, error) from error


def worksheets(workbook: Workbook) -> list[Worksheet]:
    return list(workbook.worksheets)


def sheet_rows(sheet: Worksheet) -> list[tuple[object, ...]]:
    """Rows of the used range, starting at the first used cell like calamine does."""
    rows: list[tuple[object, ...]] = [
        tuple(row)
        for row in sheet.iter_rows(
            min_row=sheet.min_row, min_col=sheet.min_column, values_only=True
        )
    ]
    if all(value is None for row in rows for value in row):
        return []
    return rows


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime | dt.date | dt.time):
        return value.isoformat()
    return str(value)
