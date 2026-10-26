"""Excel/CSV -> semantic HTML tables (header row as <th>, one <h2> per worksheet)."""

import csv
import datetime as dt
import zipfile
from html import escape
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from paddlepdf.pdftools.contract import PdfToolError

SPREADSHEET_SUFFIXES = (".csv", ".xlsx", ".xlsm", ".xltx", ".xltm")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime | dt.date | dt.time):
        return value.isoformat()
    return str(value)


def _table(rows: list[list[str]]) -> str:
    parts = ['<table repeat="1">']
    for index, row in enumerate(rows):
        tag = "th" if index == 0 else "td"
        cells = "".join(f"<{tag}>{escape(cell)}</{tag}>" for cell in row)
        parts.append(f"<tr>{cells}</tr>")
    parts.append("</table>")
    return "".join(parts)


def _csv_rows(path: Path) -> list[list[str]]:
    try:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return [list(row) for row in csv.reader(handle)]
    except (UnicodeDecodeError, csv.Error) as error:
        msg = f"Failed to parse CSV {path}: {error}"
        raise PdfToolError(msg) from error


def _workbook_sections(path: Path) -> list[tuple[str, list[list[str]]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (KeyError, ValueError, zipfile.BadZipFile, InvalidFileException) as error:
        msg = f"Failed to open workbook {path}: {error}"
        raise PdfToolError(msg) from error
    sections: list[tuple[str, list[list[str]]]] = []
    for sheet in workbook.worksheets:
        rows = [
            [_cell_text(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        sections.append((sheet.title, rows))
    workbook.close()
    return sections


def spreadsheet_html(path: Path) -> tuple[str, int]:
    """Return the HTML document and the number of data rows rendered."""
    body: list[str] = []
    row_total = 0
    if path.suffix.lower() == ".csv":
        rows = _csv_rows(path)
        body.append(_table(rows))
        row_total = max(len(rows) - 1, 0)
    else:
        for title, rows in _workbook_sections(path):
            body.append(f"<h2>{escape(title)}</h2>")
            if rows:
                body.append(_table(rows))
            row_total += max(len(rows) - 1, 0)
    return html_document("".join(body)), row_total


def html_document(body: str) -> str:
    head = '<head><meta charset="utf-8"></head>'
    return f"<!DOCTYPE html><html>{head}<body>{body}</body></html>"
